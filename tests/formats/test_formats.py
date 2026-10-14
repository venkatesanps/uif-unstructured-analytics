"""Every supported file format, end to end: readers produce labelled text, the guard reaches inside
office files and archives, classification is right, and the saved Claude answers still replay."""

import io
import json
import zipfile
from pathlib import Path

import pytest

from tests.conftest import make_ctx
from uif.contracts import GuardedDoc, NormalizedDoc, Page
from uif.evals import evaluate, gate
from uif.providers import formats as fmt
from uif.providers.local import LocalOCR
from uif.runner import process_path

ROOT = Path(__file__).resolve().parents[2]


@pytest.fixture(scope="session")
def samples(tmp_path_factory):
    from uif.synth_formats import generate_formats

    out = tmp_path_factory.mktemp("formats")
    summary = generate_formats(out)
    return out, summary


@pytest.fixture(scope="session")
def processed_samples(tmp_path_factory, samples):
    out, _ = samples
    ctx = make_ctx(tmp_path_factory.mktemp("fmt-run"))
    report = evaluate(ctx, out, run_pipeline=True)
    return ctx, report


def _texts(ctx, file_name):
    envs = [json.loads(ctx.store.get_bytes(k)) for k in ctx.store.list("envelope/")]
    doc_id = next(e["doc_id"] for e in envs if e["file_name"] == file_name)
    norm = ctx.get_model("normalized", NormalizedDoc, doc_id=doc_id)
    guarded = ctx.get_model("guarded", GuardedDoc, doc_id=doc_id)
    return norm, [p.text for p in guarded.pages]


def test_every_format_classifies_and_passes_gates(processed_samples):
    _, report = processed_samples
    assert report["documents"] >= 14
    assert report["classification_accuracy"] == 1.0
    assert gate(report) == [], report


@pytest.mark.parametrize("name,method,checks", [
    ("statement_fax.tif", "tesseract", ["Gross Sales", "Interchange Detail"]),
    ("statement_photo.png", "tesseract", ["Total Gross Sales: $"]),
    ("statement_export.xlsx", "xlsx", ["[sheet: Summary]", "Total Gross Sales: $", "[sheet: Fee Detail]"]),
    ("statement_export.csv", "csv", ["Gross Sales: $", "Credits: ($"]),
    ("rebuttal_upload.zip", "zip", ["[attachment: 1_cover_letter.docx", "[attachment: 3_delivery_photo.jpg", "Signed by"]),
    ("rebuttal_forwarded.eml", "email", ["[attachment: forwarded email: Dispute response", "[attachment: proof_of_delivery.png"]),
    ("rebuttal_portal.json", "json", ["Case Number: CB-", "Merchant Id: "]),
    ("note_memo.docx", "docx", ["Recommendation: ", "Chargeback ratio (90d): 1.9%"]),
    ("note_html_email.eml", "email", ["Case: RC-", "Merchant: "]),
    ("note_case_export.html", "html", ["Case: RC-", "Exposure estimate: $"]),
    ("note_chat.json", "json", ["Channel: merchant-risk-reviews", "Recommendation: "]),
    ("note.md", "text", ["**Recommendation:**"]),
])
def test_reader_output_keeps_labels(processed_samples, name, method, checks):
    ctx, _ = processed_samples
    norm, texts = _texts(ctx, name)
    assert method in norm.method
    joined = "\n".join(texts)
    for c in checks:
        assert c in joined, (name, c)


def test_fax_tiff_keeps_its_pages(processed_samples):
    norm, _ = _texts(processed_samples[0], "statement_fax.tif")
    assert len(norm.pages) == 2


def test_guard_reaches_inside_office_files(processed_samples):
    ctx, report = processed_samples
    _, xlsx = _texts(ctx, "statement_export.xlsx")
    _, memo = _texts(ctx, "note_memo.docx")
    assert "<PAN_" in "\n".join(xlsx) and "4111" not in "\n".join(xlsx).replace(" ", "")
    assert "<SSN_" in "\n".join(memo) and "987-65-4320" not in "\n".join(memo)
    assert report["pii_leakage"] == 0


def test_saved_claude_answers_replay_on_the_samples(samples):
    from uif import parity

    out, _ = samples
    report = parity.replay_eval(ROOT / "evals" / "formats" / "answers" / "claude-code", out)
    assert report["unanswered_documents"] == 0, "sample text drifted: regenerate requests and re-answer"
    assert gate(report) == [] and report["required_f1"] == 1.0


# ----------------------------------------------------------------------------- reader edge cases

def _zip(members: dict[str, bytes]) -> bytes:
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
        for n, d in members.items():
            z.writestr(n, d)
    return b.getvalue()


def test_zip_bomb_and_oversized_archives_are_refused():
    ocr = LocalOCR()
    with pytest.raises(NotImplementedError, match="zip bomb|expands"):
        ocr.to_pages(_zip({"zeros.txt": b"0" * 5_000_000}), "application/zip", "x.zip")
    with pytest.raises(NotImplementedError, match="limit"):
        ocr.to_pages(_zip({f"f{i}.txt": b"hello" for i in range(fmt.MAX_MEMBERS + 1)}), "application/zip", "x.zip")


def test_nested_zip_depth_and_junk_members():
    inner = _zip({"note.txt": b"Case: RC-2026-0001"})
    pages, _ = LocalOCR().to_pages(_zip({"__MACOSX/._x": b"junk", ".DS_Store": b"junk", "inner.zip": inner}),
                                   "application/zip", "outer.zip")
    assert len(pages) == 1 and "RC-2026-0001" in pages[0].text


def test_docx_page_breaks_and_tables():
    import docx

    d = docx.Document()
    d.add_paragraph("Page one text")
    d.add_paragraph().add_run().add_break(docx.enum.text.WD_BREAK.PAGE)
    d.add_paragraph("Page two text")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Exposure estimate", "$1,000.00"
    b = io.BytesIO()
    d.save(b)
    pages = fmt.docx_pages(b.getvalue())
    assert [p.page for p in pages] == [1, 2]
    assert "Exposure estimate: $1,000.00" in pages[1].text


def test_csv_dialects_and_json_flattening():
    assert "Gross Sales: $10.00" in fmt.csv_pages(b"Gross Sales;$10.00\nOther;x\n")[0].text
    text = fmt.json_pages(json.dumps({"caseNumber": "CB-2026-000001", "merchant": {"merchantId": "8000"}}).encode())[0].text
    assert "Case Number: CB-2026-000001" in text and "Merchant Id: 8000" in text


def test_audio_uses_the_transcriber_and_halts_cleanly_without_one(ctx):
    class FakeTranscriber(fmt.Transcriber):
        def transcribe(self, data, file_name):
            return "Case RC 2026 0001. Recommendation monitor."

    pages, method = LocalOCR(transcriber=FakeTranscriber()).to_pages(b"RIFF....", "audio/wav", "memo.wav")
    assert method == "transcript" and "Recommendation monitor" in pages[0].text
    unavailable = fmt.WhisperCppTranscriber(binary="/nonexistent/whisper-cli", model="/nonexistent/model.bin")
    with pytest.raises(NotImplementedError, match="whisper"):
        LocalOCR(transcriber=unavailable).to_pages(b"RIFF....", "audio/wav", "memo.wav")
    from uif.runner import process_bytes

    ctx.providers.ocr = LocalOCR(transcriber=unavailable)
    run = process_bytes(ctx, b"RIFF fake audio", "memo.wav", "internal", "test")
    last = run.results[-1]
    assert last.step == "normalize" and last.halt and "whisper" in last.detail
