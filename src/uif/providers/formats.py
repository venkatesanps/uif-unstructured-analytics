"""Readers for office files, data exports, archives and audio. Each turns bytes into pages of
text that keep the document's labels ("Total Gross Sales: $1,234.56"), so the same schemas,
hints and prompts work whatever the container format was.

Containers (ZIP, email) recurse through a `to_pages` callback, so in AWS their images and PDFs
still go to Textract while text formats are read here."""

from __future__ import annotations

import csv
import datetime as dt
import io
import json
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from typing import Any, Callable
from xml.etree import ElementTree as ET

from uif.contracts import Page

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ZIP_TYPES = {"application/zip", "application/x-zip-compressed"}

_EXTRA_MIME = {".eml": "message/rfc822", ".msg": "application/vnd.ms-outlook", ".md": "text/markdown",
               ".csv": "text/csv", ".txt": "text/plain", ".json": "application/json", ".docx": DOCX,
               ".xlsx": XLSX, ".zip": "application/zip", ".tif": "image/tiff", ".tiff": "image/tiff",
               ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png", ".wav": "audio/wav",
               ".mp3": "audio/mpeg", ".m4a": "audio/mp4", ".html": "text/html", ".htm": "text/html",
               ".pdf": "application/pdf"}


def guess_mime(file_name: str) -> str:
    import mimetypes

    ext = Path(file_name).suffix.lower()
    return _EXTRA_MIME.get(ext) or mimetypes.guess_type(file_name)[0] or "application/octet-stream"


# Archive limits: a merchant upload is a few documents, not a zip bomb.
MAX_MEMBERS, MAX_TOTAL_BYTES, MAX_RATIO, MAX_DEPTH = 50, 50 * 1024 * 1024, 100, 2


# ----------------------------------------------------------------------------- Word

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def docx_pages(data: bytes) -> list[Page]:
    """Paragraphs and tables in reading order; explicit page breaks start a new page."""
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        root = ET.fromstring(z.read("word/document.xml"))
    body = root.find(f"{_W}body")
    pages, cur = [], []

    def para_text(p) -> tuple[str, bool]:
        parts, brk = [], False
        for node in p.iter():
            if node.tag == f"{_W}t":
                parts.append(node.text or "")
            elif node.tag == f"{_W}tab":
                parts.append("\t")
            elif node.tag == f"{_W}br" and node.get(f"{_W}type") == "page":
                brk = True
        return "".join(parts), brk

    for block in list(body) if body is not None else []:
        if block.tag == f"{_W}p":
            text, brk = para_text(block)
            cur.append(text)
            if brk:
                pages.append(cur)
                cur = []
        elif block.tag == f"{_W}tbl":
            for row in block.iter(f"{_W}tr"):
                cells = [" ".join(para_text(p)[0] for p in cell.iter(f"{_W}p")).strip()
                         for cell in row.iter(f"{_W}tc")]
                cur.append(_row_text(cells))
    pages.append(cur)
    return [Page(page=i + 1, text="\n".join(p).strip()) for i, p in enumerate(pages) if "".join(p).strip()]


# ----------------------------------------------------------------------------- spreadsheets


def _row_text(cells: list[str]) -> str:
    cells = [c for c in (str(x).strip() for x in cells) if c]
    if len(cells) == 2:  # key/value layout reads like a labelled line
        return f"{cells[0]}: {cells[1]}"
    return " | ".join(cells)


def _cell_text(value: Any, number_format: str = "General") -> str:
    if value is None:
        return ""
    if isinstance(value, (dt.datetime, dt.date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        fmt = number_format or "General"
        if "$" in fmt:
            return f"${value:,.2f}"
        if "%" in fmt:
            return f"{value * 100:.2f}%"
        if "0.00" in fmt:
            return f"{value:,.2f}"
        if isinstance(value, float) and not value.is_integer():
            return f"{value:,.2f}"
        return f"{int(value):,}" if "#,##0" in fmt else str(int(value))
    return str(value)


def xlsx_pages(data: bytes) -> list[Page]:
    """One page per sheet; numbers rendered with their cell format so "$1,234.56" stays "$1,234.56"."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    pages = []
    for ws in wb.worksheets:
        lines = [f"[sheet: {ws.title}]"]
        for row in ws.iter_rows():
            text = _row_text([_cell_text(c.value, getattr(c, "number_format", "General")) for c in row])
            if text:
                lines.append(text)
        if len(lines) > 1:
            pages.append(Page(page=len(pages) + 1, text="\n".join(lines)))
    return pages


def csv_pages(data: bytes) -> list[Page]:
    text = data.decode("utf-8-sig", "replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [_row_text(r) for r in csv.reader(io.StringIO(text), dialect)]
    return [Page(page=1, text="\n".join(r for r in rows if r))]


# ----------------------------------------------------------------------------- JSON


def _humanize(key: str) -> str:
    return re.sub(r"[_\-]+", " ", re.sub(r"([a-z])([A-Z])", r"\1 \2", str(key))).strip().title()


def _flatten(obj: Any, prefix: str = "") -> list[str]:
    if isinstance(obj, dict):
        out = []
        for k, v in obj.items():
            label = f"{prefix} {_humanize(k)}".strip()
            if isinstance(v, (dict, list)):
                out.append(f"{label}:")
                out += ["  " + line for line in _flatten(v)]
            else:
                out.append(f"{_humanize(k)}: {v}")
        return out
    if isinstance(obj, list):
        out = []
        for item in obj:
            if isinstance(item, (dict, list)):
                out += _flatten(item)
                out.append("")
            else:
                out.append(f"- {item}")
        return out
    return [str(obj)]


def json_pages(data: bytes) -> list[Page]:
    """Chat exports (a "messages" list with text) read as a conversation; anything else is
    flattened to "Key: value" lines with readable keys (case_number -> Case Number)."""
    obj = json.loads(data.decode("utf-8-sig"))
    msgs = obj.get("messages") if isinstance(obj, dict) else None
    if isinstance(msgs, list) and msgs and isinstance(msgs[0], dict) and "text" in msgs[0]:
        head = [f"{_humanize(k)}: {v}" for k, v in obj.items() if k != "messages" and not isinstance(v, (dict, list))]
        lines = head + [""] + [f"[{m.get('ts', '')}] {m.get('user') or m.get('from') or 'unknown'}: {m['text']}"
                               for m in msgs]
        return [Page(page=1, text="\n".join(lines).strip())]
    return [Page(page=1, text="\n".join(_flatten(obj)).strip())]


# ----------------------------------------------------------------------------- archives


def zip_pages(data: bytes, to_pages: Callable[[bytes, str, str], tuple[list[Page], str]],
              guess_mime: Callable[[str], str], depth: int = 0) -> tuple[list[Page], str]:
    if depth >= MAX_DEPTH:
        raise NotImplementedError("nested archive too deep")
    pages, methods = [], {"zip"}
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        members = [i for i in z.infolist() if not i.is_dir() and not i.filename.startswith("__MACOSX/")
                   and not Path(i.filename).name.startswith(".")]
        if len(members) > MAX_MEMBERS:
            raise NotImplementedError(f"archive has {len(members)} files (limit {MAX_MEMBERS})")
        total = sum(i.file_size for i in members)
        packed = max(1, sum(i.compress_size for i in members))
        if total > MAX_TOTAL_BYTES or total / packed > MAX_RATIO:
            raise NotImplementedError("archive expands too much; refusing (possible zip bomb)")
        for info in sorted(members, key=lambda i: i.filename):
            name = info.filename
            mime = guess_mime(name)
            body = z.read(info)
            try:
                if mime in ZIP_TYPES:
                    sub, method = zip_pages(body, to_pages, guess_mime, depth + 1)
                else:
                    sub, method = to_pages(body, mime, Path(name).name)
            except NotImplementedError as exc:
                pages.append(Page(page=len(pages) + 1, text=f"[attachment: {name}] not read: {exc}"))
                continue
            methods.add(method)
            for p in sub:
                pages.append(Page(page=len(pages) + 1, text=f"[attachment: {name}, page {p.page}]\n{p.text}"))
    return pages, "+".join(sorted(methods))


# ----------------------------------------------------------------------------- audio


class Transcriber:
    """Speech-to-text for call recordings. Locally: whisper.cpp's CLI if installed
    (`whisper-cli` + a model at UIF_WHISPER_MODEL); in AWS: Amazon Transcribe."""

    def transcribe(self, data: bytes, file_name: str) -> str:
        raise NotImplementedError


class WhisperCppTranscriber(Transcriber):
    def __init__(self, binary: str | None = None, model: str | None = None):
        self.binary = binary or shutil.which("whisper-cli") or shutil.which("whisper-cpp")
        self.model = model or os.environ.get("UIF_WHISPER_MODEL")

    def available(self) -> bool:
        return bool(self.binary and self.model and Path(self.model).exists() and shutil.which("ffmpeg"))

    def transcribe(self, data: bytes, file_name: str) -> str:
        if not self.available():
            raise NotImplementedError("audio: install whisper.cpp (whisper-cli) and ffmpeg, and set "
                                      "UIF_WHISPER_MODEL to a ggml model file; in AWS, Transcribe is used")
        with tempfile.TemporaryDirectory() as tmp:
            src, wav = Path(tmp) / Path(file_name).name, Path(tmp) / "audio16k.wav"
            src.write_bytes(data)
            subprocess.run(["ffmpeg", "-loglevel", "error", "-y", "-i", str(src), "-ar", "16000", "-ac", "1", str(wav)],
                           check=True)
            out = subprocess.run([self.binary, "-m", self.model, "-f", str(wav), "-nt", "-np"],
                                 capture_output=True, text=True, check=True)
            return out.stdout.strip()


def audio_pages(data: bytes, file_name: str, transcriber: Transcriber) -> list[Page]:
    text = transcriber.transcribe(data, file_name)
    return [Page(page=1, text=f"[transcript of {file_name}]\n{text}")]
