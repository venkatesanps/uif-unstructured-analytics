"""Labeled samples in every common unstructured format, for testing the readers end to end.

Kept apart from the pinned golden set (synthetic/), so adding formats never shifts golden
request IDs. Same label shape as synth.py, so `uif eval samples/formats` scores them.

  merchant/  statement photo (PNG), 2-page fax (TIFF, Group 4), Excel export (XLSX, with a planted
             card number), CSV export, rebuttal upload (ZIP: Word letter + PDF invoice + JPEG photo),
             forwarded rebuttal (EML inside EML), dispute-portal API payload (JSON)
  internal/  Word memo (DOCX, planted SSN), HTML-only email, case-system export (HTML table),
             team chat export (JSON), Markdown note (audio: see synth_audio.py)
  external/  adverse-media web page (HTML), network operating bulletin (PDF): unknown types"""

from __future__ import annotations

import calendar
import datetime as dt
import io
import json
import random
import shutil
import subprocess
import tempfile
import zipfile
from email.message import EmailMessage
from pathlib import Path

from uif.synth import (
    BASIS_SENTENCE,
    ENCLOSURE,
    MERCHANTS,
    REASONS,
    THEME_SENTENCES,
    _evidence_page,
    _money,
    _note_text,
    _pdf,
    _png,
    statement_content,
)

PLANTED_PAN = "4111 1111 1111 1111"
PLANTED_SSN = "987-65-4320"


# ----------------------------------------------------------------------------- image helpers

def _photo(lines: list[str], angle: float = 1.6, fmt: str = "PNG") -> bytes:
    """A phone photo of a page: skewed, slightly blurred, grey paper."""
    from PIL import Image, ImageFilter

    img = Image.open(io.BytesIO(_png(lines))).convert("L")
    img = img.rotate(angle, expand=True, fillcolor=235).filter(ImageFilter.GaussianBlur(0.6))
    img = img.point(lambda v: int(v * 0.92 + 8))
    out = io.BytesIO()
    img.save(out, format=fmt, **({"quality": 80} if fmt == "JPEG" else {}))
    return out.getvalue()


def _fax_tiff(pages: list[list[str]]) -> bytes:
    """Multi-page bilevel TIFF (CCITT Group 4), like a faxed document."""
    from PIL import Image

    frames = [Image.open(io.BytesIO(_png(p))).convert("L").point(lambda v: 255 if v > 150 else 0).convert("1")
              for p in pages]
    out = io.BytesIO()
    frames[0].save(out, format="TIFF", save_all=True, append_images=frames[1:], compression="group4", dpi=(200, 200))
    return out.getvalue()


def _docx(paragraphs: list[str], table: list[tuple[str, str]] | None = None, title: str | None = None) -> bytes:
    import docx

    d = docx.Document()
    if title:
        d.add_heading(title, level=1)
    for line in paragraphs:
        d.add_paragraph(line)
    if table:
        t = d.add_table(rows=0, cols=2)
        for k, v in table:
            row = t.add_row().cells
            row[0].text, row[1].text = k, v
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


# ----------------------------------------------------------------------------- generator

def _note_fields(rng, merchant, analyst, date, themes, action, exposure, case_id):
    return {"case_id": case_id, "merchant_name": merchant[2], "analyst": analyst, "review_date": date,
            "risk_themes": sorted(themes), "recommended_action": action, "exposure_estimate": exposure}


def _tts_wav(text: str, dest: Path) -> bool:
    """Voice memo via macOS `say` or espeak(-ng), converted to 16 kHz WAV. False if no TTS engine."""
    with tempfile.TemporaryDirectory() as tmp:
        raw = Path(tmp) / "speech.aiff"
        if shutil.which("say"):
            cmd = ["say", "-o", str(raw), text]
        elif shutil.which("espeak-ng") or shutil.which("espeak"):
            raw = raw.with_suffix(".wav")
            cmd = [shutil.which("espeak-ng") or shutil.which("espeak"), "-w", str(raw), text]
        else:
            return False
        if not shutil.which("ffmpeg"):
            return False
        subprocess.run(cmd, check=True, capture_output=True)
        subprocess.run(["ffmpeg", "-loglevel", "error", "-y", "-i", str(raw), "-ar", "16000", "-ac", "1", str(dest)],
                       check=True)
    return True


def generate_formats(out: str | Path, seed: int = 2026) -> dict:
    rng = random.Random(seed)
    out = Path(out)
    for sub in ("merchant", "internal", "external", "labels"):
        (out / sub).mkdir(parents=True, exist_ok=True)
    labels, skipped = [], []

    def add(rel: str, data: bytes, source: str, doc_type: str, fields: dict, merchant_id, fmt: str,
            flags=(), raw_ids=()):
        (out / rel).write_bytes(data)
        labels.append({"file": rel, "source_group": source, "doc_type": doc_type, "fields": fields,
                       "merchant_id": merchant_id, "format": fmt, "flags": list(flags),
                       "raw_identifiers": [r.replace(" ", "") for r in raw_ids]})

    # 1. phone photo of a statement's first page
    m = MERCHANTS[0]
    p1, _, f = statement_content(rng, m, "A", 2026, 7, [])
    add("merchant/statement_photo.png", _photo(p1), "merchant", "processor_statement", f, m[0], "image/png (photo)")

    # 2. faxed 2-page statement
    m = MERCHANTS[1]
    p1, p2, f = statement_content(rng, m, "B", 2026, 8, [])
    add("merchant/statement_fax.tif", _fax_tiff([p1, p2]), "merchant", "processor_statement", f, m[0],
        "image/tiff (multi-page fax)", flags=["multi_page_image"], raw_ids=["disputes@lakeside-example.com"])

    # 3. Excel export with a planted card number on a notes sheet
    from openpyxl import Workbook

    m = MERCHANTS[2]
    _, _, f = statement_content(rng, m, "A", 2026, 6, [])
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Merchant Statement Export"])
    usd, cnt = '"$"#,##0.00', "#,##0"
    start, end = dt.date.fromisoformat(f["statement_period"]["start"]), dt.date.fromisoformat(f["statement_period"]["end"])
    net = f["gross_volume"] - f["refund_volume"] - f["chargeback_volume"] - f["total_fees"]
    rows = [("Processor", "Northwind Payments", None), ("Merchant Name", m[2], None),
            ("Merchant ID (MID)", m[1], "@"), ("Statement Period", f"{start:%m/%d/%Y} to {end:%m/%d/%Y}", None),
            ("Total Gross Sales", f["gross_volume"], usd), ("Total Transactions", f["transaction_count"], cnt),
            ("Refunds", f["refund_volume"], usd), ("Chargebacks", f["chargeback_count"], cnt),
            ("Chargeback Amount", f["chargeback_volume"], usd), ("Total Fees", f["total_fees"], usd),
            ("Net Deposit", round(net, 2), usd)]
    for k, v, fmt in rows:
        ws.append([k, v])
        if fmt:
            ws.cell(row=ws.max_row, column=2).number_format = fmt
    ws2 = wb.create_sheet("Fee Detail")
    for k, share in [("Interchange and assessments", 0.82), ("Processing markup", 0.15), ("Monthly service fee", 0.03)]:
        ws2.append([k, round(f["total_fees"] * share, 2)])
        ws2.cell(row=ws2.max_row, column=2).number_format = usd
    ws3 = wb.create_sheet("Notes")
    ws3.append([f"Test sale on card {PLANTED_PAN} reversed same day."])
    buf = io.BytesIO()
    wb.save(buf)
    add("merchant/statement_export.xlsx", buf.getvalue(), "merchant", "processor_statement", f, m[0],
        "xlsx (3 sheets, currency formats)", flags=["planted_pan"], raw_ids=[PLANTED_PAN])

    # 4. CSV export (layout B labels)
    m = MERCHANTS[3]
    p1, p2, f = statement_content(rng, m, "B", 2026, 5, [])
    rows = []
    for line in p1 + p2:
        line = line.lstrip("# ").strip()
        if not line:
            continue
        if ": " in line:
            rows.append(line.split(": ", 1))
        elif " $" in line:
            k, v = line.rsplit(" $", 1)
            rows.append([k, "$" + v])
        else:
            rows.append([line])
    buf = io.StringIO()
    import csv

    csv.writer(buf).writerows(rows)
    add("merchant/statement_export.csv", buf.getvalue().encode(), "merchant", "processor_statement", f, m[0], "csv",
        raw_ids=["disputes@lakeside-example.com"])

    # 5. ZIP upload: Word cover letter + PDF invoice + JPEG delivery photo
    m = MERCHANTS[4]
    sale = dt.date(2026, 8, 6)
    case, order, amount = f"CB-2026-{rng.randint(100000, 999999)}", f"ORD-{rng.randint(10000, 99999)}", 238.40
    ctx = {"order": order, "sale_date": sale, "amount": amount, "network_title": "Visa",
           "delivered": sale + dt.timedelta(days=3), "item": "Trail jacket", "signature": True}
    letter = ["Chargeback Rebuttal", f"Merchant: {m[2]}", f"MID: {m[1]}", f"Case Number: {case}",
              "Card Network: Visa", f"Reason Code: 13.1 ({REASONS[('visa', '13.1')]})",
              f"Disputed Amount: {_money(amount)}", f"Transaction Date: {sale.isoformat()}", f"Order Number: {order}",
              "", BASIS_SENTENCE["delivered"], "", "Enclosed:"] + [f"- {ENCLOSURE[k]}" for k in
                                                                  ("invoice/receipt", "proof of delivery", "signature")]
    pod = _evidence_page("proof of delivery", rng, ctx)  # PNG; re-encode as a JPEG photo
    from PIL import Image

    jpg = io.BytesIO()
    Image.open(io.BytesIO(pod)).convert("L").rotate(-1.2, expand=True, fillcolor=240).save(jpg, format="JPEG", quality=75)
    z = io.BytesIO()
    with zipfile.ZipFile(z, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("1_cover_letter.docx", _docx(letter))
        zf.writestr("2_invoice.pdf", _pdf([_evidence_page("invoice/receipt", rng, ctx)]))
        zf.writestr("3_delivery_photo.jpg", jpg.getvalue())
    add("merchant/rebuttal_upload.zip", z.getvalue(), "merchant", "dispute_rebuttal",
        {"case_number": case, "merchant_legal_name": m[2], "mid": m[1], "card_network": "visa", "reason_code": "13.1",
         "dispute_amount": amount, "transaction_date": sale.isoformat(), "order_number": order,
         "evidence_types": ["invoice/receipt", "proof of delivery", "signature"], "rebuttal_basis": "delivered"},
        m[0], "zip (docx + pdf + jpeg)", flags=["archive"])

    # 6. forwarded email: account manager forwards the merchant's portal email (EML inside EML)
    m = MERCHANTS[5]
    sale = dt.date(2026, 8, 21)
    case, order, amount = f"CB-2026-{rng.randint(100000, 999999)}", f"ORD-{rng.randint(10000, 99999)}", 412.15
    ctx = {"order": order, "sale_date": sale, "amount": amount, "network_title": "Mastercard",
           "delivered": sale + dt.timedelta(days=5), "item": "Espresso grinder", "signature": False}
    inner = EmailMessage()
    inner["From"], inner["To"] = "billing@copperleaf-example.com", "disputes@bank-example.com"
    inner["Subject"] = f"Dispute response {case}"
    inner.set_content("\n".join([
        "Chargeback dispute response submitted via merchant portal.", "Network reason code and amount are listed below.",
        f"Case Number: {case}", f"Merchant Name: {m[2]}", f"Merchant ID: {m[1]}", "Network: Mastercard",
        f"Reason: 4855 - {REASONS[('mastercard', '4855')]}", f"Amount in Dispute: {_money(amount)}",
        f"Original Sale Date: {sale:%m/%d/%Y}", f"Order #: {order}", f"Rebuttal: {BASIS_SENTENCE['delivered']}", "",
        "Attachments:", f"- {ENCLOSURE['invoice/receipt']}", f"- {ENCLOSURE['proof of delivery']}"]))
    inner.add_attachment(_pdf([_evidence_page("invoice/receipt", rng, ctx)]), maintype="application", subtype="pdf",
                         filename="invoice.pdf")
    inner.add_attachment(_evidence_page("proof of delivery", rng, ctx), maintype="image", subtype="png",
                         filename="proof_of_delivery.png")
    outer = EmailMessage()
    outer["From"], outer["To"] = "relationship.manager@bank-example.com", "merchant-risk-queue@bank-example.com"
    outer["Subject"] = f"FW: Dispute response {case}"
    outer.set_content("Forwarding the merchant's dispute response below for your review. Thanks.")
    outer.add_attachment(inner)
    add("merchant/rebuttal_forwarded.eml", bytes(outer), "merchant", "dispute_rebuttal",
        {"case_number": case, "merchant_legal_name": m[2], "mid": m[1], "card_network": "mastercard",
         "reason_code": "4855", "dispute_amount": amount, "transaction_date": sale.isoformat(), "order_number": order,
         "evidence_types": ["invoice/receipt", "proof of delivery"], "rebuttal_basis": "delivered"},
        m[0], "eml (forwarded eml with pdf + png)", flags=["nested_email"])

    # 7. dispute-portal API payload
    m = MERCHANTS[6]
    case, order, amount = f"CB-2026-{rng.randint(100000, 999999)}", f"ORD-{rng.randint(10000, 99999)}", 1310.0
    payload = {
        "type": "chargeback_response", "caseNumber": case, "submittedBy": "portal",
        "merchant": {"name": m[2], "merchantId": m[1]},
        "network": "Mastercard", "reasonCode": "4837", "disputedAmount": amount,
        "transactionDate": "2026-08-17", "orderNumber": order,
        "rebuttal": BASIS_SENTENCE["authorized"],
        "evidence": [{"type": "authorization_record", "description": ENCLOSURE["avs/cvv match"]},
                     {"type": "three_ds", "description": ENCLOSURE["3ds authentication"]},
                     {"type": "session_log", "description": ENCLOSURE["ip/device match"]}],
    }
    add("merchant/rebuttal_portal.json", json.dumps(payload, indent=2).encode(), "merchant", "dispute_rebuttal",
        {"case_number": case, "merchant_legal_name": m[2], "mid": m[1], "card_network": "mastercard",
         "reason_code": "4837", "dispute_amount": amount, "transaction_date": "2026-08-17", "order_number": order,
         "evidence_types": ["avs/cvv match", "3ds authentication", "ip/device match"], "rebuttal_basis": "authorized"},
        m[0], "json (API payload)")

    # 8-12: analyst notes in office/web/chat formats
    def note_params(i):
        merchant = MERCHANTS[(i + 2) % len(MERCHANTS)]
        themes = sorted(rng.sample(list(THEME_SENTENCES), 2))
        return (merchant, f"RC-2026-{rng.randint(1000, 9999)}", ["J. Ortiz", "P. Raman", "K. Nguyen", "S. Whitfield"][i % 4],
                (dt.date(2026, 9, 10) + dt.timedelta(days=i * 3)).isoformat(), themes,
                ["monitor", "reserve", "escalate", "reduce_limit", "terminate"][i % 5], round(rng.uniform(20_000, 300_000), -2))

    m, case_id, analyst, date, themes, action, exposure = note_params(0)
    text = _note_text(rng, m, case_id, analyst, date, themes, action, exposure, extra=f"Owner SSN {PLANTED_SSN} confirmed.")
    add("internal/note_memo.docx", _docx(text.splitlines(), table=[("Chargeback ratio (90d)", "1.9%"),
                                                                   ("Reserve held", "$0.00")], title="Risk Review Memo"),
        "internal", "risk_analyst_note", _note_fields(rng, m, analyst, date, themes, action, exposure, case_id), m[0],
        "docx (paragraphs + table)", flags=["planted_ssn"], raw_ids=[PLANTED_SSN])

    m, case_id, analyst, date, themes, action, exposure = note_params(1)
    text = _note_text(rng, m, case_id, analyst, date, themes, action, exposure)
    html_body = "<html><body style='font-family:Arial'>" + "".join(
        f"<p>{('<b>' + l.split(':', 1)[0] + ':</b>' + l.split(':', 1)[1]) if ':' in l else l}</p>"
        for l in text.splitlines() if l.strip()) + "<p>Regards,<br>Merchant Risk</p></body></html>"
    msg = EmailMessage()
    msg["From"], msg["To"] = "raman.risk@bank-example.com", "merchant-risk-queue@bank-example.com"
    msg["Subject"] = f"{case_id}: {m[2]} risk review"
    msg.set_content(html_body, subtype="html")
    add("internal/note_html_email.eml", bytes(msg), "internal", "risk_analyst_note",
        _note_fields(rng, m, analyst, date, themes, action, exposure, case_id), m[0], "eml (HTML-only body)")

    m, case_id, analyst, date, themes, action, exposure = note_params(2)
    narrative = " ".join(THEME_SENTENCES[t].format(pct=2.4, x=4) for t in themes)
    html = f"""<!doctype html><html><head><title>Case {case_id}</title><style>td{{padding:4px}}</style>
<script>window.analytics = {{}};</script></head><body><nav>Home | Cases | Reports</nav>
<h1>Risk case export</h1><table>
<tr><th>Case</th><td>{case_id}</td></tr><tr><th>Merchant</th><td>{m[2]}</td></tr>
<tr><th>Analyst</th><td>{analyst}</td></tr><tr><th>Review date</th><td>{date}</td></tr>
<tr><th>Recommendation</th><td>{action}</td></tr><tr><th>Exposure estimate</th><td>{_money(exposure)}</td></tr>
</table><h2>Risk review summary</h2><p>{narrative}</p><footer>Exported from case management</footer></body></html>"""
    add("internal/note_case_export.html", html.encode(), "internal", "risk_analyst_note",
        _note_fields(rng, m, analyst, date, themes, action, exposure, case_id), m[0], "html (table + narrative)")

    m, case_id, analyst, date, themes, action, exposure = note_params(3)
    t0 = dt.datetime.fromisoformat(date + "T09:05:00")
    opener = "R. Patel" if analyst != "R. Patel" else "S. Whitfield"  # team lead opening the case
    chat = {"channel": "merchant-risk-reviews", "exported": f"{date}T17:00:00Z", "messages": [
        {"ts": (t0).isoformat(), "user": opener, "text": f"Opening case {case_id} on {m[2]}, MID {m[1]}. Can someone take the risk review?"},
        {"ts": (t0 + dt.timedelta(minutes=12)).isoformat(), "user": analyst, "text": f"Taking it. {THEME_SENTENCES[themes[0]].format(pct=2.2, x=5)}"},
        {"ts": (t0 + dt.timedelta(minutes=40)).isoformat(), "user": analyst, "text": THEME_SENTENCES[themes[1]].format(pct=2.2, x=5)},
        {"ts": (t0 + dt.timedelta(minutes=55)).isoformat(), "user": opener, "text": "What's the analyst recommendation and exposure?"},
        {"ts": (t0 + dt.timedelta(minutes=58)).isoformat(), "user": analyst,
         "text": f"Recommendation: {action}. Exposure estimate: {_money(exposure)}. Review date {date}."},
    ]}
    add("internal/note_chat.json", json.dumps(chat, indent=2).encode(), "internal", "risk_analyst_note",
        _note_fields(rng, m, analyst, date, themes, action, exposure, case_id), m[0], "json (team chat export)")

    m, case_id, analyst, date, themes, action, exposure = note_params(4)
    narrative = " ".join(THEME_SENTENCES[t].format(pct=1.7, x=3) for t in themes)
    md = (f"# Risk review: {m[2]}\n\n| Field | Value |\n|---|---|\n| Case | {case_id} |\n| Analyst | {analyst} |\n"
          f"| Review date | {date} |\n\n## Findings\n\n{narrative}\n\n## Decision\n\n- **Recommendation:** {action}\n"
          f"- **Exposure estimate:** {_money(exposure)}\n")
    add("internal/note.md", md.encode(), "internal", "risk_analyst_note",
        _note_fields(rng, m, analyst, date, themes, action, exposure, case_id), m[0], "markdown")

    # Audio lives in its own set (samples/audio, synth_audio.py) so this set is identical on every machine.

    # 14-15. external documents of types the registry doesn't know
    article = f"""<!doctype html><html><head><title>Local news</title><script>var ads=[1,2,3]</script></head><body>
<header><a href="/">Metro Ledger</a> | Business | Weather</header><article><h1>Shoppers say online gadget orders never arrived</h1>
<p>Several customers told the paper they waited weeks for orders from {MERCHANTS[4][3]} without updates.</p>
<p>The company said a warehouse move caused backlogs and that refunds are being processed.</p></article>
<aside>Subscribe for $1 a week</aside></body></html>"""
    add("external/adverse_media.html", article.encode(), "external", "unknown", {}, None, "html (news article)")
    bulletin = [["# Card Network Operating Bulletin", "Effective date: 2027-01-15",
                 "Subject: Updated monitoring thresholds for card-absent merchants",
                 "Acquirers must notify affected merchants within 30 days.",
                 "Merchants in MCC 5967 and 5966 are in scope of the revised program.",
                 "Questions: contact your network representative."]]
    add("external/network_bulletin.pdf", _pdf(bulletin), "external", "unknown", {}, None, "pdf (bulletin)")

    for lab in labels:
        (out / "labels" / (Path(lab["file"]).name + ".json")).write_text(json.dumps(lab, indent=2))
    return {"out": str(out), "documents": len(labels), "formats": {l["file"]: l["format"] for l in labels},
            "skipped": skipped}
